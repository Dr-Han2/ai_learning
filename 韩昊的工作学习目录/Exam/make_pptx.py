"""
@Author  : 韩昊
@Company : ******
@Date    : 
@Desc    : 
@Version : 1.0
"""
#!/usr/bin/env python3
# -*- coding: utf-8 -*-
#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
根据《数据库版学生管理系统实战.md》笔记生成 PPT
依赖：pip install python-pptx
运行：python make_pptx.py
输出：学生管理系统（数据库版）实战笔记.pptx
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

# ============================ 全局配置 ============================

COLOR_PRIMARY = RGBColor(0x1F, 0x3B, 0x63)   # 深蓝
COLOR_ACCENT  = RGBColor(0x2E, 0x86, 0xDE)   # 亮蓝
COLOR_TEXT    = RGBColor(0x33, 0x33, 0x33)   # 正文
COLOR_MUTED   = RGBColor(0x6B, 0x6B, 0x6B)   # 次要文字
COLOR_CODE    = RGBColor(0x1B, 0x4B, 0x7A)   # 代码文字
COLOR_WHITE   = RGBColor(0xFF, 0xFF, 0xFF)
COLOR_SOFT    = RGBColor(0xEE, 0xF3, 0xF9)   # 表格斑马纹
COLOR_LIGHT   = RGBColor(0xBB, 0xD4, 0xF0)

FONT_CN   = "微软雅黑"
FONT_CODE = "Consolas"

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


# ============================ 基础工具 ============================

def set_run_font(run, name=FONT_CN, size=18, bold=False,
                 color=COLOR_TEXT, italic=False):
    """设置 run 的字体（同时设置中文字体 a:ea，否则中文会回退）"""
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    run.font.name = name            # 会设置 a:latin

    # 通过 font._rPr 拿到底层 <a:rPr> 元素
    rPr = getattr(run.font, "_rPr", None)
    if rPr is None:
        return

    latin = rPr.find(qn("a:latin"))
    if latin is None:
        return

    # 按 OOXML 顺序：a:latin → a:ea → a:cs
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = rPr.makeelement(qn("a:ea"), {"typeface": name})
        latin.addnext(ea)
    else:
        ea.set("typeface", name)

    cs = rPr.find(qn("a:cs"))
    if cs is None:
        cs = rPr.makeelement(qn("a:cs"), {"typeface": name})
        ea.addnext(cs)
    else:
        cs.set("typeface", name)


def add_rect(slide, left, top, width, height, color):
    """添加一个纯色矩形"""
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shp.fill.solid()
    shp.fill.fore_color.rgb = color
    shp.line.fill.background()
    try:
        shp.shadow.inherit = False
    except Exception:
        pass
    return shp


def new_slide(prs, title):
    """新建一张带标题的内容页"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_rect(slide, 0, 0, SLIDE_W, Inches(0.07), COLOR_ACCENT)
    tb = slide.shapes.add_textbox(Inches(0.55), Inches(0.26),
                                  Inches(12.2), Inches(0.68))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = title
    set_run_font(r, FONT_CN, 25, True, COLOR_PRIMARY)
    add_rect(slide, Inches(0.6), Inches(1.0), Inches(1.3), Pt(4), COLOR_ACCENT)
    return slide


def write_blocks(tf, blocks):
    """
    向文本框写入内容块
    blocks 元素格式：
        ('head', 文本)                    小标题
        ('b', 文本) / ('b', 文本, 1)      项目符号（level 0/1）
        ('code', 文本)                    等宽代码行
        ('gap', 磅值)                     空行
    """
    tf.word_wrap = True
    first = True
    for blk in blocks:
        kind = blk[0]
        if first:
            p = tf.paragraphs[0]
            first = False
        else:
            p = tf.add_paragraph()

        if kind == 'head':
            r = p.add_run()
            r.text = blk[1]
            set_run_font(r, FONT_CN, 16, True, COLOR_ACCENT)
            p.space_before = Pt(10)
            p.space_after = Pt(4)

        elif kind == 'b':
            level = blk[2] if len(blk) > 2 else 0
            prefix = "• " if level == 0 else "– "
            r = p.add_run()
            r.text = prefix + blk[1]
            if level == 0:
                set_run_font(r, FONT_CN, 14, False, COLOR_TEXT)
            else:
                set_run_font(r, FONT_CN, 12, False, COLOR_MUTED)
            p.level = level
            p.space_after = Pt(2)

        elif kind == 'code':
            r = p.add_run()
            r.text = blk[1]
            set_run_font(r, FONT_CODE, 10.5, False, COLOR_CODE)
            p.space_after = Pt(0)

        elif kind == 'gap':
            r = p.add_run()
            r.text = ""
            set_run_font(r, FONT_CN, blk[1] if len(blk) > 1 else 6)
            p.space_after = Pt(0)


def add_content_slide(prs, title, blocks, top=1.25, height=5.9):
    """标题 + 内容块"""
    slide = new_slide(prs, title)
    tb = slide.shapes.add_textbox(Inches(0.62), Inches(top),
                                  Inches(12.1), Inches(height))
    write_blocks(tb.text_frame, blocks)
    return slide


def add_table_slide(prs, title, headers, rows, col_widths,
                    top=1.35, font_size=13, row_height=0.45):
    """标题 + 表格"""
    slide = new_slide(prs, title)
    total_w = sum(col_widths)
    left = Inches((13.333 - total_w) / 2.0)
    n_rows = len(rows) + 1
    shape = slide.shapes.add_table(
        n_rows, len(headers), left, Inches(top),
        Inches(total_w), Inches(row_height * n_rows)
    )
    table = shape.table
    for i, w in enumerate(col_widths):
        table.columns[i].width = Inches(w)
    for i in range(n_rows):
        table.rows[i].height = Inches(row_height)

    for j, h in enumerate(headers):
        cell = table.cell(0, j)
        cell.text = ""
        p = cell.text_frame.paragraphs[0]
        r = p.add_run()
        r.text = h
        set_run_font(r, FONT_CN, font_size, True, COLOR_WHITE)
        cell.fill.solid()
        cell.fill.fore_color.rgb = COLOR_PRIMARY
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.margin_left = Inches(0.09)
        cell.margin_right = Inches(0.09)

    for i, row in enumerate(rows, start=1):
        for j, val in enumerate(row):
            cell = table.cell(i, j)
            cell.text = ""
            p = cell.text_frame.paragraphs[0]
            r = p.add_run()
            r.text = val
            set_run_font(r, FONT_CN, font_size, False, COLOR_TEXT)
            cell.fill.solid()
            cell.fill.fore_color.rgb = COLOR_WHITE if i % 2 else COLOR_SOFT
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.margin_left = Inches(0.09)
            cell.margin_right = Inches(0.09)
    return slide


def pit(code, symptom, cause, solution):
    """生成一条踩坑记录的内容块"""
    return [
        ('head', code),
        ('b', "现象：" + symptom, 1),
        ('b', "原因：" + cause, 1),
        ('b', "解决：" + solution, 1),
    ]


# ============================ 各页内容 ============================

def slide_cover(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_rect(slide, 0, 0, SLIDE_W, SLIDE_H, COLOR_PRIMARY)
    add_rect(slide, Inches(1.0), Inches(2.05), Inches(1.5), Pt(6), COLOR_ACCENT)

    tb = slide.shapes.add_textbox(Inches(1.0), Inches(2.35),
                                  Inches(11.3), Inches(1.6))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = "学生管理系统（数据库版）"
    set_run_font(r, FONT_CN, 40, True, COLOR_WHITE)
    p2 = tf.add_paragraph()
    r = p2.add_run()
    r.text = "项目实战笔记"
    set_run_font(r, FONT_CN, 40, True, COLOR_WHITE)

    tb2 = slide.shapes.add_textbox(Inches(1.0), Inches(4.2),
                                   Inches(11.3), Inches(0.7))
    tf2 = tb2.text_frame
    tf2.word_wrap = True
    p = tf2.paragraphs[0]
    r = p.add_run()
    r.text = "Python + SQLAlchemy + PyMySQL  |  三层架构 CLI 实战"
    set_run_font(r, FONT_CN, 18, False, COLOR_LIGHT)

    tb3 = slide.shapes.add_textbox(Inches(1.0), Inches(5.05),
                                   Inches(11.3), Inches(0.8))
    tf3 = tb3.text_frame
    tf3.word_wrap = True
    p = tf3.paragraphs[0]
    r = p.add_run()
    r.text = "架构设计 · 搭建流程 · 踩坑排查 · 最佳实践 · 扩展方向"
    set_run_font(r, FONT_CN, 14, False, RGBColor(0x8F, 0xB4, 0xDC))


def slide_toc(prs):
    add_content_slide(prs, "目录", [
        ('b', "一、项目整体逻辑 —— 分层架构 · 调用链 · 职责边界"),
        ('b', "二、项目目录结构"),
        ('b', "三、项目搭建顺序：从下往上，依赖先行"),
        ('b', "四、各层职责与模板：model / dao / service / main"),
        ('b', "五、踩过的坑 & 解决方案（A~H 共 8 类 30+ 条）"),
        ('b', "六、关键设计原则"),
        ('b', "七、快速自查清单"),
        ('b', "八、常用代码片段"),
        ('b', "九、项目可扩展方向"),
        ('b', "十、一句话总结  +  附：项目完成度记录"),
    ])


def slide_arch(prs):
    add_content_slide(prs, "一、项目整体逻辑 —— 分层架构", [
        ('code', "┌──────────────────────────────────────────────┐"),
        ('code', "│  main.py      CLI 用户交互：登录 / 菜单 / 输入输出 / 异常展示 │"),
        ('code', "└───────────────────────┬──────────────────────┘"),
        ('code', "                        │ 调用"),
        ('code', "┌───────────────────────▼──────────────────────┐"),
        ('code', "│  service/     业务逻辑：参数校验 / 唯一性校验 / 关联校验 / 填时间 │"),
        ('code', "└───────────────────────┬──────────────────────┘"),
        ('code', "                        │ 调用"),
        ('code', "┌───────────────────────▼──────────────────────┐"),
        ('code', "│  dao/         数据访问：拼 SQL / execute / commit / rollback │"),
        ('code', "└───────────────────────┬──────────────────────┘"),
        ('code', "                        │ 依赖"),
        ('code', "┌───────────────────────▼──────────────────────┐"),
        ('code', "│  model/       数据模型：@dataclass 纯数据类 + from_dict       │"),
        ('code', "└──────────────────────────────────────────────┘"),
        ('gap', 6),
        ('b', "横切关注点：utils/db_util（连接池 + Session）· utils/env_util（环境变量）"),
        ('b', "　　　　　　core/logger（日志）· core/exceptions（异常体系）"),
    ])


def slide_call_chain(prs):
    add_content_slide(prs, "一次完整调用链（以「新增学生」为例）", [
        ('code', "main.py  student_add()"),
        ('code', "  │ 收集输入 → Student 对象"),
        ('code', "  ▼"),
        ('code', "StudentService.add(student)"),
        ('code', "  │ 1. 必填校验（姓名、学号）"),
        ('code', "  │ 2. 学号唯一性校验  → StudentDao.select_by_stu_num"),
        ('code', "  │ 3. 班级存在性校验  → ClazzDao.select_by_id"),
        ('code', "  │ 4. 年龄 / 手机号格式校验"),
        ('code', "  │ 5. 填 create_time / update_time"),
        ('code', "  ▼"),
        ('code', "StudentDao.insert(student)"),
        ('code', "  │ session = DBUtil.get_session()       ← 从连接池拿"),
        ('code', "  │ session.execute(text(sql), params)   ← 参数绑定"),
        ('code', "  │ session.commit()                     ← 落库"),
        ('code', "  │ student.id = result.lastrowid        ← 回填主键"),
        ('code', "  │ finally: session.close()             ← 归还连接"),
        ('code', "  ▼"),
        ('code', "回到 main，打印 student.id"),
    ])


def slide_duty(prs):
    add_table_slide(
        prs,
        "职责边界：每一层只做自己该做的事",
        ["层", "该做", "不该做"],
        [
            ["model", "定义字段、from_dict", "任何数据库 / SQL"],
            ["dao", "拼 SQL、execute、commit/rollback/close、异常包装", "业务校验、填时间"],
            ["service", "参数校验、唯一性 / 关联校验、填时间、调多个 dao", "直接拼 SQL、管理 session"],
            ["main", "输入输出、异常展示", "业务逻辑"],
            ["utils/core", "横切工具（连接池、环境变量、日志、异常）", "具体业务"],
        ],
        col_widths=[1.5, 6.6, 3.4],
        font_size=13,
        row_height=0.62,
    )


def slide_tree(prs):
    add_content_slide(prs, "二、项目目录结构", [
        ('code', "05_学生管理系统_数据库/"),
        ('code', "├── .env                       # 环境变量（放在项目根目录）"),
        ('code', "├── main.py                    # CLI 入口"),
        ('code', "├── README.md"),
        ('code', "│"),
        ('code', "├── core/                      # 核心工具"),
        ('code', "│   ├── exceptions.py          # 异常体系"),
        ('code', "│   └── logger.py              # 日志（依赖 env_util）"),
        ('code', "│"),
        ('code', "├── utils/                     # 通用工具"),
        ('code', "│   ├── db_util.py             # 数据库工具（连接池 + session）"),
        ('code', "│   └── env_util.py            # 环境变量工具"),
        ('code', "│"),
        ('code', "├── model/                     # 数据模型（纯数据类）"),
        ('code', "│   ├── clazz.py  student.py  course.py  score.py  user.py"),
        ('code', "│"),
        ('code', "├── dao/                       # 数据访问"),
        ('code', "│   └── clazz_dao.py  student_dao.py  course_dao.py ..."),
        ('code', "│"),
        ('code', "├── service/                   # 业务逻辑"),
        ('code', "│   └── clazz_service.py  student_service.py ..."),
        ('code', "│"),
        ('code', "└── logs/                      # 日志目录"),
        ('code', "    └── student.log"),
    ])


def slide_build_order(prs):
    add_content_slide(prs, "三、项目搭建顺序", [
        ('code', "1. 设计数据库表 → 建库 SQL"),
        ('code', "2. 写 .env                → 数据库 + 日志配置"),
        ('code', "3. 写 utils/env_util      → 读环境变量"),
        ('code', "4. 写 core/exceptions     → 定义异常体系"),
        ('code', "5. 写 core/logger         → 日志（依赖 env_util）"),
        ('code', "6. 写 utils/db_util       → 连接池 + get_session"),
        ('code', "7. 每张表："),
        ('code', "     7.1 写 model/X.py"),
        ('code', "     7.2 写 dao/X_dao.py"),
        ('code', "     7.3 写 service/X_service.py（先 add + get_by_id，跑通）"),
        ('code', "     7.4 补全其他方法"),
        ('code', "8. 最后写 main.py 串起来"),
        ('gap', 12),
        ('head', "核心原则"),
        ('b', "从下往上写，依赖先行。上层调下层，下层不依赖上层。"),
    ])


def slide_model(prs):
    add_content_slide(prs, "四、model 层 —— 纯数据类", [
        ('code', "from dataclasses import dataclass"),
        ('code', "from datetime import datetime"),
        ('code', ""),
        ('code', "@dataclass"),
        ('code', "class Student:"),
        ('code', '    """学生类"""'),
        ('code', "    id: int | None = None"),
        ('code', "    stu_num: str | None = None"),
        ('code', "    name: str | None = None"),
        ('code', "    gender: str | None = None"),
        ('code', "    age: int | None = None"),
        ('code', "    class_id: int | None = None"),
        ('code', "    phone: str | None = None"),
        ('code', "    create_time: datetime | None = None"),
        ('code', "    update_time: datetime | None = None"),
        ('code', ""),
        ('code', "    @staticmethod"),
        ('code', "    def from_dict(dic):"),
        ('code', '        """从字典到对象"""'),
        ('code', "        if dic is None:"),
        ('code', "            return None"),
        ('code', "        return Student(**dic)"),
        ('gap', 6),
        ('b', "要点：所有字段默认 None · 类型注解用 X | None · from_dict 统一处理 None 输入"),
    ])


def slide_dao(prs):
    add_content_slide(prs, "四、dao 层模板 —— 数据访问", [
        ('code', "class StudentDao:"),
        ('code', '    _COLUMNS = ("`id`, `stu_num`, `name`, `gender`, `age`, "'),
        ('code', '                "`class_id`, `phone`, `create_time`, `update_time`")'),
        ('code', ""),
        ('code', "    @staticmethod"),
        ('code', "    def insert(student: Student) -> int:"),
        ('code', "        session = None"),
        ('code', "        try:"),
        ('code', '            sql = ("INSERT INTO `t_student`(...) "'),
        ('code', '                   "VALUES (:stu_num, :name, :gender, ...)")'),
        ('code', "            params = {...}"),
        ('code', "            session = DBUtil.get_session()"),
        ('code', "            result = session.execute(text(sql), params)"),
        ('code', "            session.commit()                 # 关键：提交事务"),
        ('code', "            student.id = result.lastrowid    # 回填自增主键"),
        ('code', "            return result.rowcount           # 返回受影响行数"),
        ('code', "        except Exception as e:"),
        ('code', "            if session:"),
        ('code', "                session.rollback()           # 失败回滚"),
        ('code', '            log.exception("插入学生记录失败")  # 自动带堆栈'),
        ('code', "            raise DatabaseError(e) from e    # 保留原因链"),
        ('code', "        finally:"),
        ('code', "            if session:"),
        ('code', "                session.close()              # 归还连接池"),
    ])


def slide_service(prs):
    add_content_slide(prs, "四、service 层模板 —— 业务逻辑", [
        ('code', "class StudentService:"),
        ('code', '    """学生业务逻辑层"""'),
        ('code', ""),
        ('code', "    @staticmethod"),
        ('code', "    def add(student: Student) -> int:"),
        ('code', "        # 1. 必填校验"),
        ('code', "        if student is None:"),
        ('code', '            raise ValidationError("学生信息不能为空")'),
        ('code', "        if not student.name or not student.name.strip():"),
        ('code', '            raise ValidationError("学生姓名不能为空", field="name")'),
        ('code', "        # 2. 唯一性校验"),
        ('code', "        if StudentDao.select_by_stu_num(student.stu_num) is not None:"),
        ('code', '            raise AlreadyExistsError("学生", field="stu_num", value=...)'),
        ('code', "        # 3. 关联存在性校验"),
        ('code', "        if student.class_id is not None:"),
        ('code', "            if ClazzDao.select_by_id(student.class_id) is None:"),
        ('code', '                raise NotFoundError("班级", student.class_id)'),
        ('code', "        # 4. 时间字段"),
        ('code', "        now = datetime.now()"),
        ('code', "        student.create_time = now"),
        ('code', "        student.update_time = now"),
        ('code', "        # 5. 调用 DAO"),
        ('code', "        row = StudentDao.insert(student)"),
        ('code', "        if row <= 0:"),
        ('code', '            raise BusinessException("新增学生记录失败")'),
        ('code', '        log.info("新增学生成功 id=%s stu_num=%s", student.id, student.stu_num)'),
        ('code', "        return row"),
    ])


def slide_main(prs):
    add_content_slide(prs, "四、main 层模板 —— CLI 交互", [
        ('code', "def student_add():"),
        ('code', '    """新增学生"""'),
        ('code', "    student = Student()"),
        ('code', '    student.stu_num = input_str("学号: ")'),
        ('code', '    student.name = input_str("姓名: ")'),
        ('code', '    student.age = input_int("年龄(可空): ", allow_empty=True)'),
        ('code', '    student.class_id = input_int("班级ID(可空): ", allow_empty=True)'),
        ('code', "    StudentService.add(student)"),
        ('code', '    print(f"新增学生成功，ID={student.id}")'),
        ('code', ""),
        ('code', ""),
        ('code', "def student_menu():"),
        ('code', "    while True:"),
        ('code', '        print("1. 新增  2. 查询  3. 更新  4. 删除  0. 返回")'),
        ('code', "        try:"),
        ('code', '            if choice == "1":'),
        ('code', "                student_add()"),
        ('code', "        except AppException as e:"),
        ('code', '            print(f"[业务错误] {e.message}")'),
        ('code', "        except ValueError as e:"),
        ('code', '            print(f"[输入错误] {e}")'),
        ('code', "        except Exception as e:"),
        ('code', '            print(f"[系统错误] {e}")'),
        ('code', "        pause()"),
    ])


def slide_pit_a(prs):
    blocks = [('b', "A. 环境配置类（5 条）", 0)]
    blocks += pit("A1 · int() argument must be ... not 'NoneType'",
                  'DB_PORT = int(os.getenv("DB_PORT")) 启动即报错',
                  ".env 里缺这一项，或 load_dotenv() 没找到 .env",
                  '确认 .env 在运行目录；getenv 加默认值 "3306"')
    blocks += pit("A2 · key 名字对不上",
                  ".env 写 DB_NAME，代码读 DB_DATABASE → 读到 None",
                  "命名约定不统一",
                  "统一 key 名；env_util 加 _require()，读不到直接抛错")
    blocks += pit("A3 · Unknown database 'none'",
                  "SQLAlchemy 连库报 1049 Unknown database 'none'",
                  "库名读成 None，f-string 拼进 URL 变成 \"None\"",
                  "key 名对齐 + MySQL 里 CREATE DATABASE 建库，SQL 存档")
    blocks += pit("A4 · 日志文件变成 student.log.log",
                  ".env 写 LOG_NAME=student.log，文件名多一个 .log",
                  "logger.py 里会拼 .log，LOG_NAME 又带了一次",
                  "LOG_NAME=student，扩展名由代码统一加")
    blocks += pit("A5 · Level not an integer or a valid string: None",
                  "logger.setLevel(None) 报错",
                  ".env 里没配 LOG_LEVEL",
                  'LOG_LEVEL = os.getenv("LOG_LEVEL", "INFO")')
    add_content_slide(prs, "五、踩过的坑 —— A. 环境配置类", blocks)


def slide_pit_b(prs):
    blocks = [('b', "B. 数据库连接 / 会话类（5 条）", 0)]
    blocks += pit("B1 · thread_local 缓存了已关闭的 session",
                  "第一次调用正常，第二次拿到坏 session，报 Session is closed",
                  "get_session 用 thread_local 缓存，dao close 后引用还在",
                  "每次返回全新 session，不做缓存（连接池已吃掉复用收益）")
    blocks += pit("B2 · 忘记 session.close() 导致连接泄漏",
                  "程序跑久了连接耗尽，MySQL 报 Too many connections",
                  "某个 dao 方法异常时没走到 close",
                  "所有 dao 方法 try / except / finally，close 放 finally")
    blocks += pit("B3 · 忘记 session.commit() 数据没落库",
                  "日志显示「插入成功」，但数据库里查不到",
                  "execute 后没 commit，程序退出时事务回滚",
                  "DML 方法执行后必须 commit；或统一用 DAO 模板")
    blocks += pit("B4 · 重复 commit",
                  "with DBUtil.session_dml() 里调用 dao，dao 内部也 commit",
                  "事务边界不清晰 —— 两处都在管",
                  "二选一：DAO 管事务（推荐） 或 Service 用 @transactional")
    blocks += pit("B5 · 事务不原子",
                  "一次业务多次写操作，前一次成功、后一次失败，数据不一致",
                  "每次 dao 调用各自 commit，没有统一事务",
                  "@transactional 装饰 service 方法，所有 dao 共用一个 session")
    add_content_slide(prs, "五、踩过的坑 —— B. 数据库连接 / 会话类", blocks)


def slide_pit_c1(prs):
    blocks = [('b', "C. SQLAlchemy / SQL 类（上）", 0)]
    blocks += pit("C1 · 'CursorResult' object has no attribute 'inserted_id'",
                  "insert 后取主键报错",
                  "inserted_id 是 PyMongo 的 API，SQLAlchemy 里没有",
                  "MySQL 用 result.lastrowid，或 inserted_primary_key[0]")
    blocks += pit("C2 · 返回 lastrowid 而不是 rowcount",
                  "service 判断 row <= 0 失败",
                  "insert 返回的是主键 ID（可能大于 0），不是受影响行数",
                  "insert 最后 return result.rowcount，返回值语义统一")
    blocks += pit("C3 · SQL 里 : param 中间有空格",
                  "UPDATE ... SET update_time=: update_time 参数绑定失败",
                  "SQLAlchemy 参数名必须紧跟冒号，不能有空格",
                  "写成 :update_time，或让 IDE 高亮帮你检查")
    add_content_slide(prs, "五、踩过的坑 —— C. SQLAlchemy / SQL 类（上）", blocks)


def slide_pit_c2(prs):
    blocks = [('b', "C. SQLAlchemy / SQL 类（下）", 0)]
    blocks += pit("C4 · 动态 SQL 参数缩进错位",
                  "SQL 里没这个字段，params 里却有 → 参数不匹配报错",
                  "params 赋值缩进跑到 if 外面，被无条件执行",
                  "SQL 拼接和 params 赋值成对写在一起，缩进保持一致")
    blocks += pit("C5 · _COLUMNS 未定义",
                  "SELECT {XxxDao._COLUMNS} 报 NameError",
                  "类属性没定义",
                  "类顶部加 _COLUMNS 常量，所有查询统一用它")
    blocks += pit('C6 · LIKE 模糊查询误当唯一性校验',
                  '新增课程 "Python" 时，因库里已有 "Python程序设计" 报重复',
                  "select_by_name 用 LIKE，只适合搜索，不该用于唯一性校验",
                  "另写 select_by_name_exact 用 = 精确匹配，命名上区分开")
    add_content_slide(prs, "五、踩过的坑 —— C. SQLAlchemy / SQL 类（下）", blocks)


def slide_pit_de(prs):
    blocks = [('b', "D. 日志类", 0)]
    blocks += pit('D1 · log.error("xxx", e) 导致 Logging error',
                  "TypeError: not all arguments converted during string formatting",
                  "logging 第二个参数是格式化参数，消息里没 %s 却传了参数",
                  '用 log.exception("插入失败") 或 log.error("插入失败: %s", e)')
    blocks += [('gap', 8), ('b', "E. 异常处理类", 0)]
    blocks += pit("E1 · raise DatabaseError(e) 丢失原始堆栈",
                  "只能看到 DatabaseError，原始 OperationalError 被吞",
                  "没加 from e",
                  "所有异常包装都写 raise DatabaseError(e) from e")
    blocks += pit("E2 · 异常类型不统一",
                  "clazz_service 用 ValidationError，student_service 用 ValueError",
                  "没定异常体系",
                  "定义 core.exceptions，Service 抛业务异常，main 统一 catch")
    blocks += pit("E3 · AlreadyExistsError / ValidationError 继承错位",
                  "NotFoundError 是 BusinessException 子类，AlreadyExistsError 却挂在 AppException 下",
                  "异常层次定义时不严谨",
                  "所有业务异常都继承 BusinessException，先画继承树再写")
    add_content_slide(prs, "五、踩过的坑 —— D. 日志类 & E. 异常处理类", blocks)


def slide_pit_f(prs):
    blocks = [('b', "F. 业务校验类（4 条）", 0)]
    blocks += pit("F1 · 唯一性校验只在 add 里做，update 忘了",
                  "更新课程名时改成和别的课同名，居然能成功",
                  "update 缺唯一性校验",
                  "update 里也检查唯一性，并排除自己：exist.id != current.id")
    blocks += pit("F2 · 关联存在性校验只在 add 里做",
                  "更新成绩把 stu_id 改成不存在的值，能成功",
                  "update 没校验外键存在",
                  "add / update 共用同一套校验函数，外键字段全部校验")
    blocks += pit("F3 · class_id = 0 被当成有效值",
                  "用户输入 0，0 is not None 为 True，被写入库",
                  "CLI 层把空输入转成了 0",
                  "CLI 空 → None；Service 层校验 class_id <= 0 直接报错")
    blocks += pit('F4 · 字段类型和范围没校验',
                  'age 输成 999，手机号 "abc" 都能写进去',
                  "Service 缺校验",
                  "抽出 _validate_age / _validate_phone，add 和 update 都调")
    add_content_slide(prs, "五、踩过的坑 —— F. 业务校验类", blocks)


def slide_pit_g(prs):
    blocks = [('b', "G. Python / 项目结构类（4 条）", 0)]
    blocks += pit("G1 · 文件被误覆盖",
                  "model/course.py 内容变成 CourseDao",
                  "多文件操作时复制粘贴串了",
                  "检查文件内容是否与文件名相符；文件头加明显 docstring")
    blocks += pit("G2 · 文件名用 . 而不是 _",
                  "service/student.service.py",
                  "命名不规范",
                  "重命名为 student_service.py；模块名只用小写 + 下划线")
    blocks += pit("G3 · 循环导入",
                  "from model.course import Course 写在 model/course.py 里",
                  "文件名和 import 路径对不上",
                  "确认 import 路径正确；层级清晰，model 不 import dao")
    blocks += pit("G4 · 文件头 docstring 重复",
                  '顶部两个 """...""" 连着',
                  "复制模板时没注意",
                  "删掉一个，保持文件头干净")
    add_content_slide(prs, "五、踩过的坑 —— G. Python / 项目结构类", blocks)


def slide_pit_h(prs):
    blocks = [('b', "H. CLI 交互类（3 条）", 0)]
    blocks += pit("H1 · 空输入变成 0",
                  "用户直接回车，class_id 变成 0 写进数据库",
                  "input_int 没区分「空」和「0」",
                  "allow_empty=True 时空输入返回 None；0 不作为有效 ID")
    blocks += pit("H2 · int() 抛 ValueError 飘到最外层",
                  "用户输入字母，程序报 [系统错误] invalid literal for int()",
                  "输入转换没有统一收口",
                  "统一走 input_int 工具函数，main 里 except ValueError 给友好提示")
    blocks += pit("H3 · 更新时无法清空字段",
                  "update 时输入空回车 → None → DAO 跳过更新 → 字段不会变 NULL",
                  "用 None 表示「不修改」，就没办法表示「清空」",
                  "约定特殊符号（如输入 - 表示清空），设计 CLI 时想清楚")
    add_content_slide(prs, "五、踩过的坑 —— H. CLI 交互类", blocks)


def slide_principles(prs):
    add_content_slide(prs, "六、关键设计原则", [
        ('head', "6.1 单向依赖"),
        ('b', "main → service → dao → model；core / utils 作为横切被各层引用"),
        ('b', "下层不知道上层存在：model 不 import dao，dao 不 import service"),
        ('head', "6.2 事务边界清晰"),
        ('b', "推荐：DAO 管事务（自己 commit / rollback / close）"),
        ('b', "进阶：Service 用 @transactional 管，DAO 只 execute　｜　禁忌：两处都管"),
        ('head', "6.3 异常分层"),
        ('b', "AppException → BusinessException（NotFound / AlreadyExists / Validation）"),
        ('b', "　　　　　　　→ SystemException（DatabaseError / ConfigError）"),
        ('b', "Service 抛 BusinessException，DAO 抛 DatabaseError，main 统一 catch AppException"),
        ('head', "6.4 参数绑定，杜绝 SQL 注入"),
        ('b', '永远用 :param + 字典传参：session.execute(text("... = :id"), {"id": x})'),
        ('head', "6.5 校验一致 & 6.6 日志规范"),
        ('b', "add / update 共用校验函数；唯一性校验排除自己；空串与 None 统一处理"),
        ('b', 'DAO 用 log.exception 带堆栈；Service 用 log.info("xxx成功 id=%s", id)；CLI 只 print'),
    ])


def slide_checklist(prs):
    add_content_slide(prs, "七、快速自查清单", [
        ('head', "环境配置"),
        ('b', ".env 在项目根目录 · key 与 getenv 一致 · 关键变量有默认值或启动校验 · 数据库已创建"),
        ('head', "DAO 层"),
        ('b', "每个 DML 都有 commit · 每个方法都有 try/except/finally · finally 里 close"),
        ('b', "insert 返回 rowcount，主键用 lastrowid 回填 · 所有查询有 _COLUMNS 常量"),
        ('head', "Service 层"),
        ('b', "add / update 都有必填校验 · update 有存在性 + 唯一性校验（排除自己）"),
        ('b', "关联字段（外键）有存在性校验 · 时间字段自动填 · 异常用 core.exceptions 里的类型"),
        ('head', "异常"),
        ('b', "raise XxxError(e) from e 保留原因链 · main 里 except AppException 统一处理"),
        ('b', "业务异常不打印堆栈，系统异常打印"),
        ('head', "日志"),
        ('b', "log.error 用 %s 或 log.exception · LOG_NAME 不带 .log · LOG_LEVEL 有默认值"),
        ('head', "CLI"),
        ('b', "输入转换工具函数统一处理空值和类型错误 · 关键操作有确认 · 空输入转 None 而不是 0"),
    ])


def slide_snippet_1(prs):
    add_content_slide(prs, "八、常用代码片段（1）", [
        ('head', "8.1 环境变量工具 utils/env_util.py"),
        ('code', "load_dotenv()"),
        ('code', ""),
        ('code', "def _require(key):"),
        ('code', "    val = os.getenv(key)"),
        ('code', "    if not val:"),
        ('code', '        raise RuntimeError(f"环境变量 {key} 未设置")'),
        ('code', "    return val"),
        ('code', ""),
        ('code', "class EnvUtil:"),
        ('code', '    DB_HOST = _require("DB_HOST")'),
        ('code', '    DB_PORT = int(_require("DB_PORT"))'),
        ('code', '    DB_USER = _require("DB_USER")'),
        ('code', '    DB_PASSWORD = _require("DB_PASSWORD")'),
        ('code', '    DB_DATABASE = _require("DB_DATABASE")'),
        ('code', '    DB_CHARSET = os.getenv("DB_CHARSET", "utf8mb4")'),
        ('code', '    LOG_LEVEL = os.getenv("LOG_LEVEL", "INFO")'),
        ('code', '    LOG_DIR = os.getenv("LOG_DIR", "logs")'),
        ('code', '    LOG_NAME = os.getenv("LOG_NAME", "student")'),
        ('gap', 8),
        ('head', "8.3 事务装饰器（进阶：让 Service 管事务）"),
        ('code', "def transactional(func):"),
        ('code', "    @wraps(func)"),
        ('code', "    def wrapper(*args, **kwargs):"),
        ('code', "        session = DBUtil.get_session()"),
        ('code', "        try:"),
        ('code', "            result = func(*args, **kwargs)"),
        ('code', "            session.commit()"),
        ('code', "            return result"),
        ('code', "        except Exception:"),
        ('code', "            session.rollback()"),
        ('code', "            raise"),
        ('code', "        finally:"),
        ('code', "            session.close()"),
        ('code', "    return wrapper"),
    ])


def slide_snippet_2(prs):
    add_content_slide(prs, "八、常用代码片段（2）", [
        ('head', "8.2 数据库工具 utils/db_util.py"),
        ('code', "class DBUtil:"),
        ('code', "    __engine = None"),
        ('code', "    __SessionLocal = None"),
        ('code', "    __lock = threading.Lock()"),
        ('code', ""),
        ('code', "    @classmethod"),
        ('code', "    def _init_engine(cls):"),
        ('code', "        if cls.__engine is None:"),
        ('code', "            with cls.__lock:"),
        ('code', "                if cls.__engine is None:"),
        ('code', '                    url = (f"mysql+pymysql://{EnvUtil.DB_USER}:{EnvUtil.DB_PASSWORD}"'),
        ('code', '                           f"@{EnvUtil.DB_HOST}:{EnvUtil.DB_PORT}/{EnvUtil.DB_DATABASE}"'),
        ('code', '                           f"?charset={EnvUtil.DB_CHARSET}")'),
        ('code', "                    cls.__engine = create_engine("),
        ('code', "                        url, pool_size=5, max_overflow=10, pool_timeout=30,"),
        ('code', "                        pool_recycle=3600, pool_pre_ping=True, echo=False)"),
        ('code', "                    cls.__SessionLocal = sessionmaker("),
        ('code', "                        bind=cls.__engine, autoflush=False, autocommit=False)"),
        ('code', ""),
        ('code', "    @classmethod"),
        ('code', "    def get_session(cls) -> Session:"),
        ('code', "        cls._init_engine()"),
        ('code', "        return cls.__SessionLocal()   # 每次新建，不做缓存"),
        ('gap', 8),
        ('head', "8.7 建库 SQL 要点"),
        ('code', "CREATE DATABASE student_manage_db CHARACTER SET utf8mb4;"),
        ('code', "ALTER TABLE `t_class` ADD UNIQUE INDEX `uk_class_name`(`name`);  -- 唯一索引兜底"),
    ])


def slide_extend(prs):
    add_content_slide(prs, "九、项目可扩展方向", [
        ('b', "9.1 事务进阶 —— 引入 @transactional，Service 管事务，支持跨多 DAO 的原子操作（如学生转班）"),
        ('b', "9.2 密码加密 —— t_user.password 目前是明文，生产环境用 bcrypt / passlib"),
        ('code', "        hashed = bcrypt.hashpw(password.encode(), bcrypt.gensalt())"),
        ('code', "        bcrypt.checkpw(input_pw.encode(), hashed)"),
        ('b', "9.3 分页查询 —— select_all 数据量大时慢，加 LIMIT / OFFSET"),
        ('code', '        sql = "SELECT ... LIMIT :size OFFSET :offset"   # offset = (page-1)*size'),
        ('b', "9.4 唯一索引兜底 —— 应用层校验 + 数据库 UNIQUE 索引，报 Duplicate entry 由 DatabaseError 捕获"),
        ('b', "9.5 数据软删除 —— 加 is_deleted 字段，查询带 WHERE is_deleted = 0，保留历史数据"),
        ('b', "9.6 Web 化 —— main.py 换成 Flask / FastAPI，Service 层不用改，异常 to_dict() 返回 JSON，加 JWT"),
        ('b', "9.7 缓存 —— 热点数据（课程列表）用 Redis，先查缓存，miss 再查库并回写"),
        ('b', "9.8 单元测试 —— pytest + unittest.mock"),
        ('code', "        def test_student_add_duplicate_stu_num():"),
        ('code', "            with pytest.raises(AlreadyExistsError):"),
        ('code', '                StudentService.add(Student(stu_num="S001", name="张三"))'),
    ])


def slide_summary(prs):
    slide = new_slide(prs, "十、一句话总结")
    box = add_rect(slide, Inches(0.9), Inches(1.7),
                   Inches(11.5), Inches(2.5), COLOR_SOFT)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.4)
    tf.margin_right = Inches(0.4)
    tf.margin_top = Inches(0.3)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    r = p.add_run()
    r.text = ("下层稳定、上层灵活；事务边界清晰、异常分层明确；\n"
              "参数绑定防注入、校验一致防漏网；\n"
              "日志规范、DAO 模板化，剩下的就是重复劳动。")
    set_run_font(r, FONT_CN, 22, True, COLOR_PRIMARY)

    tb = slide.shapes.add_textbox(Inches(0.9), Inches(4.5),
                                  Inches(11.5), Inches(2.0))
    tf2 = tb.text_frame
    tf2.word_wrap = True
    p = tf2.paragraphs[0]
    r = p.add_run()
    r.text = "按这个思路走，三层架构的项目基本不会出大问题。"
    set_run_font(r, FONT_CN, 16, False, COLOR_TEXT)
    p2 = tf2.add_paragraph()
    r = p2.add_run()
    r.text = "新功能只需在对应的层加方法，不用动其他层。"
    set_run_font(r, FONT_CN, 16, False, COLOR_TEXT)


def slide_progress(prs):
    add_table_slide(
        prs,
        "附：项目完成度记录",
        ["模块", "Model", "DAO", "Service", "main 菜单", "状态"],
        [
            ["班级 Clazz", "✓", "✓ 6 方法", "✓ 6 方法", "✓", "✅ 完成"],
            ["学生 Student", "✓", "✓ 7 方法", "✓ 7 方法", "✓", "✅ 完成"],
            ["课程 Course", "✓", "✓ 6 方法", "✓ 5 方法", "✓", "✅ 完成"],
            ["成绩 Score", "✓", "✓ 7 方法", "✓ 7 方法", "✓", "✅ 完成"],
            ["用户 User", "✓", "✓ 1 方法", "✓ 1 方法", "✓", "✅ 完成"],
            ["日志 logger", "—", "—", "—", "—", "✅ 完成"],
            ["异常 exceptions", "—", "—", "—", "—", "✅ 完成"],
            ["数据库工具 DBUtil", "—", "—", "—", "—", "✅ 完成"],
            ["环境变量 EnvUtil", "—", "—", "—", "—", "✅ 完成"],
        ],
        col_widths=[2.6, 1.5, 2.2, 2.2, 1.8, 1.6],
        top=1.5,
        font_size=12,
        row_height=0.46,
    )
    slide = prs.slides[-1]
    tb = slide.shapes.add_textbox(Inches(0.9), Inches(6.3),
                                  Inches(11.5), Inches(0.8))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = ("项目全部功能已跑通：登录 → 班级管理 → 学生管理 → 课程管理 → 成绩管理，"
              "所有增删改查和边界校验都正常。")
    set_run_font(r, FONT_CN, 13, True, COLOR_ACCENT)


# ============================ 主流程 ============================

def main():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    slide_cover(prs)
    slide_toc(prs)
    slide_arch(prs)
    slide_call_chain(prs)
    slide_duty(prs)
    slide_tree(prs)
    slide_build_order(prs)
    slide_model(prs)
    slide_dao(prs)
    slide_service(prs)
    slide_main(prs)
    slide_pit_a(prs)
    slide_pit_b(prs)
    slide_pit_c1(prs)
    slide_pit_c2(prs)
    slide_pit_de(prs)
    slide_pit_f(prs)
    slide_pit_g(prs)
    slide_pit_h(prs)
    slide_principles(prs)
    slide_checklist(prs)
    slide_snippet_1(prs)
    slide_snippet_2(prs)
    slide_extend(prs)
    slide_summary(prs)
    slide_progress(prs)

    out = "学生管理系统（数据库版）实战笔记.pptx"
    prs.save(out)
    print(f"已生成：{out}，共 {len(prs.slides)} 页")


if __name__ == "__main__":
    main()